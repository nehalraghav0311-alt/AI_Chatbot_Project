from pptx import Presentation
import os

def extract_text_from_ppt(ppt_path, output_txt_path):
    """Extracts text and table data from a PPTX file and saves it to a text file."""
    try:
        prs = Presentation(ppt_path)
    except Exception as e:
        print(f"Error loading presentation: {e}")
        return

    with open(output_txt_path, "w", encoding="utf-8") as f:
        for slide_num, slide in enumerate(prs.slides, start=1):
            f.write(f"\n{'='*20} Slide {slide_num} {'='*20}\n\n")
            
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    f.write(shape.text.strip() + "\n\n")
                
                if shape.has_table:
                    for row in shape.table.rows:
                        row_data = []
                        for cell in row.cells:
                            cell_text = cell.text_frame.text.replace('\n', ' ').strip()
                            row_data.append(cell_text)
                        f.write(" | ".join(row_data) + "\n")
                    f.write("\n")

if __name__ == "__main__":
    INPUT_FILE = r"C:\AI_PROJECT\PPT_AI_PROJECT.pptx" 
    OUTPUT_FILE = r"C:\AI_PROJECT\extracted_knowledge_base.txt"
    
    if os.path.exists(INPUT_FILE):
        print(f"Extracting text from {INPUT_FILE}...")
        extract_text_from_ppt(INPUT_FILE, OUTPUT_FILE)
        print(f"Extraction complete. Check {OUTPUT_FILE}.")
    else:
        print(f"Error: {INPUT_FILE} not found. Please check the filename.")